# -*- coding: utf-8 -*-
"""슬라이드 다이어트: 제목 + 핵심 메시지 1줄(18pt 강조박스) + 표/이미지만 남긴다.
설명 문단·◈ 소제목·각주는 삭제(대본으로 이관). 표 글자 12pt, 긴 셀 축약.
사용: python3 ppt_diet.py <in.pptx> <out.pptx>
"""
import sys
from pptx import Presentation
from pptx.util import Pt, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

IN, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(IN)
GREEN, WHITE, DARK = RGBColor(0x1C,0x4F,0x35), RGBColor(0xFF,0xFF,0xFF), RGBColor(0x22,0x22,0x22)
LEFT, WIDTH = Cm(1.3), Cm(22.9)
MSG_TOP, MSG_H = Cm(2.7), Cm(1.15)
FONT='맑은 고딕'

def cm(v): return round(Emu(v).inches*2.54,2)

def style_table(tbl, size=12, hdr=13):
    for ri,row in enumerate(tbl.rows):
        for cell in row.cells:
            cell.margin_left=Cm(0.15); cell.margin_right=Cm(0.1); cell.margin_top=Cm(0.03); cell.margin_bottom=Cm(0.03)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.size=Pt(hdr if ri==0 else size)
                for r in p.runs:
                    r.font.size=Pt(hdr if ri==0 else size); r.font.name=FONT
                    if ri==0: r.font.bold=True

def set_cell(cell, text, size=12, bold=False, color=None, align=None):
    tf=cell.text_frame
    for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
    p=tf.paragraphs[0]
    for r in list(p.runs): r._r.getparent().remove(r._r)
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=bold; r.font.name=FONT
    if color: r.font.color.rgb=color
    if align: p.alignment=align
    p.font.size=Pt(size)

def add_table(slide, rows, top, height, colw=None, size=12, first_col_bold=True):
    n, m = len(rows), len(rows[0])
    gf=slide.shapes.add_table(n, m, LEFT, top, WIDTH, height); t=gf.table
    if colw:
        for i,w in enumerate(colw): t.columns[i].width=Cm(w)
    for ri in range(n): t.rows[ri].height=Emu(int(height/n))
    for ri,row in enumerate(rows):
        for ci,val in enumerate(row):
            c=t.cell(ri,ci)
            hdr = ri==0
            set_cell(c, val, size=(size+1 if hdr else size), bold=(hdr or (first_col_bold and ci==0)),
                     color=(WHITE if hdr else None),
                     align=(PP_ALIGN.CENTER if (hdr or ci==0) else PP_ALIGN.LEFT))
    style_table(t, size, size+1)
    return gf

def add_msg(slide, text):
    box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LEFT, MSG_TOP, WIDTH, MSG_H)
    box.adjustments[0]=0.18
    box.fill.solid(); box.fill.fore_color.rgb=RGBColor(0xEB,0xF3,0xEE)
    box.line.color.rgb=GREEN; box.line.width=Pt(1.25)
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=Cm(0.4); tf.margin_top=Cm(0.05); tf.margin_bottom=Cm(0.05)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=text; r.font.size=Pt(18); r.font.bold=True; r.font.color.rgb=GREEN; r.font.name=FONT
    return box

def set_text_shape(sh, lines, size, bold=False, top=None, height=None):
    tf=sh.text_frame; tf.word_wrap=True
    tf.clear()
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        r=p.add_run(); r.text=ln; r.font.size=Pt(size); r.font.bold=bold; r.font.name=FONT; p.space_after=Pt(4)
    if top is not None: sh.top=top
    if height is not None: sh.height=height

# ---------------- 슬라이드별 스펙 ----------------
# del: 삭제할 shape index / edit: {(shape,row,col): text} / geom: {shape:(top_cm,height_cm)}
# text: {shape: (lines, size, bold, top_cm, height_cm)} / add: [(rows, top_cm, height_cm, colw)]
S = {
2: dict(msg="군산 개정면 3,000평 고설 수경재배 — 3년차 매출 4.19억",
        dele=[2,4],
        edit={(3,2,1):"전북 군산시 개정면 (가족 농지, 부지 확보)",
              (3,3,1):"시설 3,600평 → 재배 3,000평 (16.67% 공제)",
              (3,4,1):"촉성재배: 9월 정식 → 11월~5월 수확",
              (3,6,1):"도매시장 + 대형마트 · 계약재배 + 직거래",
              (3,7,1):"2026년 설치 → 2027~2031년 5개년"},
        geom={3:(4.2,8.0), 5:(12.9,2.4)},
        colw={3:[5.0,17.9]}),
3: dict(msg="모든 숫자는 공공기관 원자료에서 — 출처 3요소 병기",
        dele=[2,3,4],
        geom={5:(4.2,10.5)},
        edit={(5,1,2):"표준단수 3,292kg/10a · 노동 755.2h/10a",
              (5,2,2):"농업노임 · 농기계 내용연수",
              (5,3,2):"설향 상품 월별 도매가격",
              (5,4,2):"재배면적 −1.7% · 설향 비중 80.2%",
              (5,5,2):"작부체계 시기 · 재식밀도",
              (5,6,2):"자가노동 2,400시간/년",
              (5,7,2):"온실 형식 선정"},
        colw={5:[9.5,6.4,7.0]}),
4: dict(msg="딸기 사업 자산 = 농지 4.68억 + 자기자금 4억, 본인 단독 경영",
        dele=[2,4,6],
        edit={(3,1,3):"평당 13만원 (국토부 실거래 2025)",
              (3,2,3):"조부 벼농사 (참고 표기)",
              (3,3,3):"조부 소유 → 미계상",
              (3,4,3):"④ 투자 자부담 전액",
              (5,1,4):"⑪ 자가 2,400h와 일치",
              (5,2,4):"토지만 확보, 노동 미계상",
              (5,3,4):"계절 고용 5,089h"},
        geom={3:(4.2,5.4), 5:(10.4,4.6)},
        colw={3:[4.6,7.6,3.6,7.1], 5:[3.4,5.5,3.4,4.6,6.0]}),
5: dict(msg="5년 목표: 매출 3.4억 → 4.2억 · 감당 가능한 규모인지 검증",
        dele=[2,4,5],
        edit={(3,2,6):"복합환경제어 도입", (3,4,6):"3차 ⑮에서 확정"},
        geom={3:(4.2,5.6)},
        colw={3:[3.6,2.7,2.7,2.7,2.7,2.7,5.8]},
        add=[([("3년차 소득","소득률","상환능력 DSCR"),
               ("5,886만원","14.0%","1.82  (원리금 7,249만 vs 영업현금 1억 3,219만)")], 11.0, 2.6, [6.0,4.5,12.4])]),
6: dict(msg="1년 농사 달력: 9월 심고 → 11월~5월 수확 · 연 755시간",
        dele=[2,3,4],
        geom={6:(4.2,7.4), 5:(12.4,4.8)}),
7: dict(msg="생산량 = 표준단수 × 재배면적 × 달성률 × 배분율 — 전부 셀 참조",
        dele=[2,4,6,7],
        text={3:(["3,292kg/10a  ×  9,917㎡  ×  90→100→110%  =  생산량",
                  "생산량  ×  상품화율 90%  =  판매량"], 15, True, 4.2, 2.3),
              5:(["⑦ 판매계획 C21:   ='8. 생산계획'!$C$10 * $C$11/1000 * $C$17 * $C$26"], 13, False, 6.7, 1.0)},
        geom={8:(8.0,9.6)},
        edit={(8,1,2):"농촌진흥청 2024 소득자료집 (시설딸기 수경)",
              (8,2,2):"시설 11,901㎡ × (1−16.67%) — 수식",
              (8,3,2):"농진청 융합모형: 3.3㎡당 12kg = 기존 온실 평균",
              (8,4,2):"자료집 99.4% 대비 보수 적용"},
        colw={8:[4.6,7.8,10.5]}),
8: dict(msg="1~2월 성출하기에 40% 집중 — 3년차 판매 32,320kg · 4.19억",
        dele=[2,4,6],
        geom={3:(4.2,3.6), 5:(8.6,5.2)}),
9: dict(msg="5년 평년가격 12,971원/kg — 최고가가 아니라 평균값을 썼다",
        dele=[2,4,5,6],
        geom={3:(4.2,10.5)}),
10: dict(msg="KAMIS에서 직접 조회 — 설향 · 상품 · kg 환산 · 2021~2025",
        dele=[2,4,5,7],
        edit={(3,1,1):"KAMIS → 가격정보 → 중도매인 판매가격 → 월간",
              (3,2,1):"채소류 / 딸기 / 설향 / 상품 / kg 환산",
              (3,3,1):"2021~2025년 월간 (연도별 4개년씩 2회 조회)"},
        geom={3:(4.2,3.8)},
        img={6:(8.4,10.0,22.9)}),
11: dict(msg="조회 결과 화면 ① — 2022~2025년 (9장 표와 동일 숫자)", dele=[3], img={2:(4.2,13.8,22.9)}),
12: dict(msg="조회 결과 화면 ② — 2021년 열 확보 (2018~20년은 미사용)", dele=[3], img={2:(4.2,13.8,22.9)}),
13: dict(msg="연 7,490시간 필요 — 본인 2,400h + 고용 5,089h (68%)",
        dele=[2,4,5,6],
        edit={(7,1,2):"시급 환산 + 농업노임 인상률", (7,2,2):"딸기(수경) 실측단가 × 3%/년"},
        geom={3:(4.2,4.8), 7:(10.0,4.0)},
        colw={7:[5.5,9.0,8.4]}),
14: dict(msg="내재해형 단동 온실 — 설치비 31만원/평 (연동 46만원)",
        dele=[2,3,4,6],
        edit={(5,1,1):"내재해형 단동 비닐하우스 다동 배치",
              (5,1,2):"설치비 단동 31만 vs 연동 46만/평, 수량 차이 작음",
              (5,3,2):"광열비 절감 (239만원/10a)",
              (5,4,2):"온 · 습도 · CO₂ · 양액 정밀관리"},
        geom={5:(4.2,7.2)},
        colw={5:[4.0,8.5,10.4]}),
15: dict(msg="교수님 1차 검토기준 10개 — 모두 미리 점검했다", geom={2:(4.2,13.6)}),
16: dict(msg="3년차 매출 4.19억 · 소득 5,886만원 · DSCR 1.82",
        dele=[2,4,5],
        edit={(3,1,1):"농지 + 자기자금 4억",
              (3,2,1):"매출 3.4→4.2억 · 소득 14%",
              (3,3,1):"촉성 1기작 · 작업별 노동시간",
              (3,4,1):"단수 × 면적 × 달성률 — 전 인자 수식",
              (3,5,1):"5개년 평년가격 · 방식 1",
              (3,6,1):"자가 2,400h + 고용 5,089h",
              (3,1,2):"국토부 실거래가 2025",
              (3,2,2):"표준단수 · 평년가격",
              (3,3,2):"농사로 표준 작업일정",
              (3,4,2):"2024 소득자료집",
              (3,5,2):"KAMIS · 농업관측 2026.9",
              (3,6,2):"NCS · 소득자료집"},
        geom={3:(4.2,7.6)},
        colw={3:[4.2,10.5,8.2]},
        add=[([("다음 단계","할 일"),
               ("2차","모델농장 섭외 · 군산 내재해 기준 · 온실 견적"),
               ("3차","⑮ 손익계획 완성 → ② 순이익 갱신")], 12.4, 3.4, [4.0,18.9])]),
}

for sl, spec in S.items():
    slide = prs.slides[sl-1]
    shapes = list(slide.shapes)
    # 1) 삭제
    for i in sorted(spec.get('dele',[]), reverse=True):
        el=shapes[i]._element; el.getparent().remove(el)
    # 2) 메시지 박스
    add_msg(slide, spec['msg'])
    # 3) 텍스트 편집 (원 index 기준)
    for i,(lines,size,bold,top,h) in spec.get('text',{}).items():
        set_text_shape(shapes[i], lines, size, bold, Cm(top), Cm(h))
    # 4) 표 셀 편집
    for (i,r,c),txt in spec.get('edit',{}).items():
        set_cell(shapes[i].table.cell(r,c), txt, 12)
    # 5) 표 기하 + 스타일
    for i,(top,h) in spec.get('geom',{}).items():
        sh=shapes[i]; sh.left=LEFT; sh.width=WIDTH; sh.top=Cm(top); sh.height=Cm(h)
        if sh.has_table:
            t=sh.table; n=len(t.rows)
            for row in t.rows: row.height=Emu(int(Cm(h)/n))
            if i in spec.get('colw',{}):
                for ci,w in enumerate(spec['colw'][i]): t.columns[ci].width=Cm(w)
            style_table(t, 12, 13)
    # 6) 이미지 기하 (비율 유지, 폭 기준)
    for i,(top,h,w) in spec.get('img',{}).items():
        sh=shapes[i]; ratio=sh.height/sh.width
        sh.left=LEFT; sh.top=Cm(top); sh.width=Cm(w); sh.height=int(Cm(w)*ratio)
        if sh.height>Cm(h): sh.height=Cm(h); sh.width=int(Cm(h)/ratio); sh.left=int(LEFT+(WIDTH-sh.width)/2)
    # 7) 새 표
    for rows,top,h,colw in spec.get('add',[]):
        add_table(slide, rows, Cm(top), Cm(h), colw)
    print(f"slide {sl}: ok")

prs.save(OUT); print("saved:", OUT)

# ---------------- 검증 ----------------
chk=Presentation(OUT); H=cm(chk.slide_height)
issues=[]
for si,s in enumerate(chk.slides,1):
    chars=0
    for sh in s.shapes:
        bottom=cm(sh.top+sh.height)
        if bottom>H-0.3: issues.append(f"s{si} 하단 넘침 {bottom}cm")
        if sh.has_text_frame and not sh.has_table:
            chars+=len(sh.text_frame.text)
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size<Pt(12): issues.append(f"s{si} {r.font.size.pt}pt")
    print(f"  s{si}: 텍스트박스 글자수 {chars}")
print("이슈:", issues if issues else "없음")
