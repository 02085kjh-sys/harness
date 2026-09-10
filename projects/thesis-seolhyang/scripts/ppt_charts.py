# -*- coding: utf-8 -*-
"""숫자 슬라이드를 [차트 + KPI 카드] 스타일로 변환 (2·5·8·9·13장)"""
import sys
from pptx import Presentation
from pptx.util import Pt, Cm, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION

IN, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(IN)
GREEN=RGBColor(0x1C,0x4F,0x35); GOLD=RGBColor(0xB8,0x86,0x0B); GRAY=RGBColor(0x66,0x66,0x66)
LIGHT=RGBColor(0xF3,0xF6,0xF4); DARK=RGBColor(0x22,0x22,0x22); FONT='맑은 고딕'
LEFT=Cm(1.3); WIDTH=Cm(22.9)

def tables(s): return [sh for sh in s.shapes if sh.has_table]
def kill(sh): sh._element.getparent().remove(sh._element)

def card(slide, x, y, w, h, label, value, unit="", accent=False, vsize=26):
    box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.adjustments[0]=0.10; box.fill.solid(); box.fill.fore_color.rgb=LIGHT; box.line.fill.background()
    box.shadow.inherit=False
    tf=box.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=Cm(0.5); tf.margin_right=Cm(0.3); tf.margin_top=Cm(0.15); tf.margin_bottom=Cm(0.15)
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
    r=p.add_run(); r.text=label; r.font.size=Pt(12); r.font.color.rgb=GRAY; r.font.name=FONT
    p2=tf.add_paragraph(); p2.alignment=PP_ALIGN.LEFT; p2.space_before=Pt(2)
    r2=p2.add_run(); r2.text=value; r2.font.size=Pt(vsize); r2.font.bold=True; r2.font.name=FONT
    r2.font.color.rgb=GOLD if accent else GREEN
    if unit:
        r3=p2.add_run(); r3.text=" "+unit; r3.font.size=Pt(14); r3.font.color.rgb=GRAY; r3.font.name=FONT
    return box

def note(slide, text, y=Cm(17.6)):
    tb=slide.shapes.add_textbox(LEFT, y, WIDTH, Cm(0.9)); tf=tb.text_frame; tf.word_wrap=True
    r=tf.paragraphs[0].add_run(); r.text="* "+text; r.font.size=Pt(12); r.font.color.rgb=GRAY; r.font.name=FONT
    return tb

def bar(slide, x, y, w, h, cats, vals, title=None, fmt='#,##0', color=GREEN, label_size=12, chart_type=XL_CHART_TYPE.COLUMN_CLUSTERED):
    cd=CategoryChartData(); cd.categories=cats; cd.add_series('', vals)
    gf=slide.shapes.add_chart(chart_type, x, y, w, h, cd); ch=gf.chart
    ch.has_legend=False
    if title:
        ch.has_title=True; ch.chart_title.text_frame.text=title
        tp=ch.chart_title.text_frame.paragraphs[0]; tp.runs[0].font.size=Pt(14); tp.runs[0].font.bold=True; tp.runs[0].font.name=FONT
    else: ch.has_title=False
    pl=ch.plots[0]; pl.gap_width=60; pl.has_data_labels=True
    dl=pl.data_labels; dl.number_format=fmt; dl.number_format_is_linked=False
    dl.font.size=Pt(label_size); dl.font.bold=True; dl.font.name=FONT; dl.position=XL_LABEL_POSITION.OUTSIDE_END
    ser=pl.series[0]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb=color
    ca=ch.category_axis; ca.tick_labels.font.size=Pt(12); ca.tick_labels.font.name=FONT
    ca.format.line.color.rgb=RGBColor(0xBB,0xBB,0xBB); ca.has_major_gridlines=False
    va=ch.value_axis; va.has_major_gridlines=True; va.major_gridlines.format.line.color.rgb=RGBColor(0xE3,0xE3,0xE3)
    va.tick_labels.font.size=Pt(10); va.tick_labels.font.name=FONT; va.tick_labels.number_format=fmt; va.tick_labels.number_format_is_linked=False
    va.format.line.fill.background()
    return gf

# ---------- slide 2: KPI 표 → 4장 카드 ----------
s=prs.slides[1]
for t in tables(s):
    if len(t.table.rows)==2: kill(t)
big=tables(s)[0]; big.top=Cm(4.2); big.height=Cm(8.6)
for r in big.table.rows: r.height=Emu(int(Cm(8.6)/8))
cw=Cm(5.45); gap=Cm(0.37)
for i,(lab,val,unit) in enumerate([("3년차 생산량","35,911","kg"),("3년차 판매량","32,320","kg"),("3년차 매출액","4.19","억원"),("가중평균단가","12,971","원/kg")]):
    card(s, LEFT+i*(cw+gap), Cm(13.4), cw, Cm(3.6), lab, val, unit, accent=(i==2), vsize=24)
print("slide 2 ok")

# ---------- slide 5: 표 2개 → 연도별 매출 차트 + 카드 3장 ----------
s=prs.slides[4]
for t in tables(s): kill(t)
bar(s, LEFT, Cm(4.1), Cm(14.6), Cm(12.6),
    ["2027\n(90%)","2028\n(100%)","2029\n(110%)","2030\n(110%)","2031\n(110%)"],
    [342994,381104,419215,419215,419215], title="연도별 매출액 (천원) — 괄호는 목표달성률")
cx=Cm(16.4); cwid=Cm(7.8)
card(s, cx, Cm(4.1), cwid, Cm(3.9), "3년차 소득 (내 손에 남는 돈)", "5,886", "만원", accent=True, vsize=28)
card(s, cx, Cm(8.3), cwid, Cm(3.9), "소득률", "14.0", "%", vsize=28)
card(s, cx, Cm(12.5), cwid, Cm(4.2), "상환능력 DSCR", "1.82", "배", vsize=28)
note(s, "재배면적 9,917㎡ 5년 고정 · 순이익은 3차 ⑮ 손익계획 확정 후 갱신 · DSCR = 영업현금 1억 3,219만 ÷ 원리금 7,249만")
print("slide 5 ok")

# ---------- slide 8: 배분표 → 월별 생산량 차트, 연도표는 우측 ----------
s=prs.slides[7]
for t in tables(s):
    if len(t.table.columns)==9: kill(t)
yt=tables(s)[0]
bar(s, LEFT, Cm(4.1), Cm(14.0), Cm(12.8),
    ["1월","2월","3월","4월","5월","6월","7월","8월","9월","10월","11월","12월"],
    [5876,5876,5289,4113,2351,0,0,0,0,0,1469,4407], title="2027년 월별 생산량 (kg)")
yt.left=Cm(15.8); yt.top=Cm(4.1); yt.width=Cm(8.4); yt.height=Cm(6.4)
for r in yt.table.rows: r.height=Emu(int(Cm(6.4)/4))
for ci,w in enumerate([2.4,2.0,2.0,2.0]): yt.table.columns[ci].width=Cm(w)
yt.table.cell(0,1).text_frame.paragraphs[0].runs[0].text="2027"; yt.table.cell(0,2).text_frame.paragraphs[0].runs[0].text="2028"; yt.table.cell(0,3).text_frame.paragraphs[0].runs[0].text="2029~"
for row in yt.table.rows:
    for c in row.cells:
        for p in c.text_frame.paragraphs:
            for r in p.runs: r.font.size=Pt(11)
card(s, Cm(15.8), Cm(11.0), Cm(8.4), Cm(2.8), "1~2월 성출하기 집중", "40", "%", accent=True, vsize=24)
card(s, Cm(15.8), Cm(14.1), Cm(8.4), Cm(2.8), "3년차 판매량", "32,320", "kg", vsize=24)
note(s, "6~10월은 여름철 휴지기(육묘·정식기)로 출하 없음 · 배분율은 농사로 화방별 수확곡선 기준, 합계 100%")
print("slide 8 ok")

# ---------- slide 9: KAMIS 표 → 평년가격 차트 + 카드 ----------
s=prs.slides[8]
for t in tables(s): kill(t)
bar(s, LEFT, Cm(4.1), Cm(14.6), Cm(12.6),
    ["11월","12월","1월","2월","3월","4월","5월"],
    [14120,18414,16433,13234,9975,8262,7712], title="설향 상품 5개년 평년가격 (원/kg, 출하 순서)")
cx=Cm(16.4); cwid=Cm(7.8)
card(s, cx, Cm(4.1), cwid, Cm(4.0), "가중평균단가 (판매량 가중)", "12,971", "원/kg", accent=True, vsize=28)
card(s, cx, Cm(8.4), cwid, Cm(4.0), "최고 단가 — 12월", "18,414", "원/kg", vsize=26)
card(s, cx, Cm(12.7), cwid, Cm(4.0), "최저 단가 — 5월", "7,712", "원/kg", vsize=26)
note(s, "KAMIS 중도매인 판매가격 · 2021~2025년 평균 · 연도별 원자료 표는 11·12장 조회 화면 참조 · 결측월은 가용연도 평균 대체")
print("slide 9 ok")

# ---------- slide 13: 노동 배분 차트 + 카드, 노임표 축소 ----------
s=prs.slides[12]
ts=tables(s); lab=[t for t in ts if len(t.table.rows)==4][0]; wage=[t for t in ts if len(t.table.rows)==3][0]
kill(lab)
bar(s, LEFT, Cm(4.1), Cm(12.4), Cm(8.2),
    ["본인(자가)","고용 남","고용 여"], [2400,875,4214], title="연간 노동 배분 (시간)")
cx=Cm(14.2); cwid=Cm(10.0)
card(s, cx, Cm(4.1), cwid, Cm(3.9), "연간 필요 노동", "7,490", "시간  (755.2h/10a × 9.917)", accent=True, vsize=28)
card(s, cx, Cm(8.3), cwid, Cm(3.9), "고용 비중", "68", "%  (5,089시간)", vsize=28)
wage.left=LEFT; wage.top=Cm(12.8); wage.width=WIDTH; wage.height=Cm(4.4)
for r in wage.table.rows: r.height=Emu(int(Cm(4.4)/3))
note(s, "자가노동 상한 2,400h = NCS 기준 · 고용단가는 2024 소득자료집 딸기(수경) 실측 지급단가 (남 13,995원/h · 여 12,602원/h)")
print("slide 13 ok")

prs.save(OUT); print("saved", OUT)
chk=Presentation(OUT); H=Emu(chk.slide_height).inches*2.54
bad=[f"s{i} {round(Emu(sh.top+sh.height).inches*2.54,1)}" for i,s in enumerate(chk.slides,1) for sh in s.shapes if Emu(sh.top+sh.height).inches*2.54>H-0.2 and i!=1]
print("넘침:", bad or "없음")
print("차트 수:", sum(1 for s in chk.slides for sh in s.shapes if sh.has_chart))
